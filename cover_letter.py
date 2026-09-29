# ============================================================
# cover_letter.py
# Builds a tailored cover letter (DOCX + PDF) from resume_data.py,
# job_details.py and job_description.txt. Uses only real content.
# ============================================================

import re
import datetime
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from resume_data import PERSONAL, CERTS_SHORT, EXPERIENCE, PROJECTS
from job_details import COMPANY, ROLE, HIRING_MANAGER, WHY_COMPANY
from tailor_resume import load_jd, reorder_bullets, reorder_projects, kw_in, display_url
from pdf_export import export_pdf

UNUSED_SIEM_VENDORS = ["qradar", "arcsight", "logrhythm", "rsa"]


def as_clause(bullet_text):
    """'Triaged real-time alerts...' -> 'triaged real-time alerts...'"""
    text = bullet_text.strip().rstrip(".")
    return text[0].lower() + text[1:]


def build_letter_text(jd_lower):
    job1, job2 = EXPERIENCE[0], EXPERIENCE[1]
    b1 = reorder_bullets(job1["bullets"], jd_lower)
    b2 = reorder_bullets(job2["bullets"], jd_lower)
    top_project = reorder_projects(PROJECTS, jd_lower)[0]

    paragraphs = [
        f"I am applying for the {ROLE} position at {COMPANY}. I bring hands-on "
        "experience in enterprise multi-tenant SOC operations, including alert triage, "
        "incident investigation and first-response containment using Microsoft Sentinel "
        "and Microsoft Defender XDR.",

        f"In my role as {job1['title']} at {job1['company']}, I {as_clause(b1[0]['text'])}. "
        f"I also {as_clause(b1[1]['text'])}. Earlier, at {job2['company']}, "
        f"I {as_clause(b2[0]['text'])}.",

        f"In my {top_project['name']} project, I {as_clause(top_project['bullets'][0]['text'])}. "
        f"I hold the {CERTS_SHORT} certifications.",
    ]

    if any(kw_in(v, jd_lower) for v in UNUSED_SIEM_VENDORS):
        paragraphs.append(
            "My SIEM experience is primarily with Microsoft Sentinel, with lab exposure to "
            "Splunk and Wazuh. The core workflow of log analysis, alert correlation, triage "
            "and escalation carries across platforms, and I am ready to ramp up quickly on "
            "the tools your team uses."
        )

    paragraphs.append(
        f"{WHY_COMPANY.strip()} I would welcome the opportunity to discuss how I can "
        "contribute to your security operations team. Thank you for your time and consideration."
    )
    return paragraphs


def build_docx(paragraphs):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(20)
    sec.left_margin = sec.right_margin = Mm(22)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.1

    def line(text, size=11, bold=False, color=None, center=False, after=8):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
        if color:
            r.font.color.rgb = RGBColor(*color)

    # Header (matches resume style)
    line(PERSONAL["name"], 18, True, (31, 78, 121), True, 1)
    line(f"{PERSONAL['location']} | {PERSONAL['phone']} | {PERSONAL['email']}", 9.8, center=True, after=0)
    line(f"{display_url(PERSONAL['linkedin'])} | {display_url(PERSONAL['github'])} | "
         f"{display_url(PERSONAL['portfolio'])}", 9.8, center=True, after=14)

    line(datetime.date.today().strftime("%d %B %Y"))
    line(f"{HIRING_MANAGER}\n{COMPANY}")
    line(f"Dear {HIRING_MANAGER},")
    for para in paragraphs:
        line(para)
    line("Sincerely,", after=2)
    line(PERSONAL["name"].title())
    return doc


def main():
    if not WHY_COMPANY.strip():
        raise SystemExit("Write WHY_COMPANY in job_details.py before generating the letter.")

    jd_lower = load_jd("job_description.txt").lower()
    paragraphs = build_letter_text(jd_lower)

    safe_company = re.sub(r"[^A-Za-z0-9]+", "_", COMPANY).strip("_")
    out = f"Khader_Shareef_Cover_Letter_{safe_company}.docx"
    build_docx(paragraphs).save(out)
    print(f"Saved DOCX: {out}")
    pdf = export_pdf(out)
    if pdf:
        print(f"Saved PDF:  {pdf}")
    print("Read it fully before sending, and edit any sentence that doesn't sound like you.")


if __name__ == "__main__":
    main()